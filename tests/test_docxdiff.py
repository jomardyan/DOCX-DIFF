import sys
import json
import queue
import threading
import time
from pathlib import Path
from unittest.mock import MagicMock, patch, mock_open
import pytest
import difflib
import tkinter as tk
from docx import Document

# Import functions to test from docxdiff
from docxdiff import (
    Config,
    atomic_write_text,
    build_source_metadata,
    evaluate_policy,
    sha256_file,
    validate_file,
    validate_output_targets,
    verify_source_metadata,
    iter_block_text,
    load_docx_lines,
    print_unified_diff,
    print_colored_diff,
    calculate_diff_stats,
    export_to_html,
    export_to_json,
    print_side_by_side,
    print_statistics,
    main,
    DocxDiffGUI,
)
from docxdiff_engine import (
    ComparisonCancelled,
    DocumentBlock,
    compare_blocks,
    compare_documents,
    compare_word_spans,
    extract_document_blocks,
)


def test_sha256_and_source_metadata(tmp_path):
    source = tmp_path / "contract.docx"
    source.write_bytes(b"enterprise-document")

    assert sha256_file(source) == (
        "3d276a64386730d33952a9df78e96c01ce79b108ad4b1f5d8aad80fddbd78cb6"
    )

    metadata = build_source_metadata(source, redact_path=True)
    assert metadata["path"] == "contract.docx"
    assert metadata["size_bytes"] == len(b"enterprise-document")
    assert metadata["sha256"] == sha256_file(source)


def test_verify_source_metadata_detects_changes(tmp_path):
    source = tmp_path / "contract.docx"
    source.write_bytes(b"version-one")
    metadata = build_source_metadata(source)
    source.write_bytes(b"version-two")

    with pytest.raises(RuntimeError, match="changed during comparison"):
        verify_source_metadata(source, metadata)


def test_atomic_write_text_replaces_existing_file(tmp_path):
    output = tmp_path / "report.json"
    output.write_text("old", encoding="utf-8")

    atomic_write_text(output, "new report")

    assert output.read_text(encoding="utf-8") == "new report"
    assert list(tmp_path.glob("*.tmp")) == []


def test_validate_output_targets_rejects_collisions(tmp_path):
    source = tmp_path / "source.docx"
    report = tmp_path / "report.json"

    valid, message = validate_output_targets([source], [source])
    assert not valid
    assert "overwrite an input" in message

    valid, message = validate_output_targets([report, report], [source])
    assert not valid
    assert "same path" in message


def test_evaluate_policy():
    stats = {"similarity": 75.0, "additions": 3, "deletions": 2}

    violations = evaluate_policy(stats, min_similarity=80.0, max_changes=4)

    assert len(violations) == 2
    assert "below required" in violations[0]
    assert "exceeds allowed" in violations[1]


def test_validate_file_not_found(tmp_path):
    non_existent = tmp_path / "does_not_exist.docx"
    is_valid, msg = validate_file(non_existent)
    assert not is_valid
    assert "File not found" in msg


def test_validate_file_not_a_file(tmp_path):
    is_valid, msg = validate_file(tmp_path)
    assert not is_valid
    assert "Not a file" in msg


def test_validate_file_unsupported_extension(tmp_path):
    invalid_file = tmp_path / "test.txt"
    invalid_file.write_text("hello")
    is_valid, msg = validate_file(invalid_file)
    assert not is_valid
    assert "Unsupported file type" in msg


def test_validate_file_too_large(tmp_path, monkeypatch):
    large_file = tmp_path / "large.docx"
    large_file.write_text("fake docx")

    # Mock Config.MAX_FILE_SIZE_MB to be extremely small for testing
    monkeypatch.setattr(Config, "MAX_FILE_SIZE_MB", 0.000001)

    is_valid, msg = validate_file(large_file)
    assert not is_valid
    assert "File too large" in msg


def test_validate_file_valid(tmp_path):
    valid_file = tmp_path / "valid.docx"
    valid_file.write_text("fake docx")
    is_valid, msg = validate_file(valid_file)
    assert is_valid
    assert msg == ""


def test_iter_block_text():
    # Mock a docx Document
    mock_doc = MagicMock()

    # Mock paragraphs
    mock_p1 = MagicMock()
    mock_p1.text = "  Paragraph 1  "
    mock_p2 = MagicMock()
    mock_p2.text = None
    mock_doc.paragraphs = [mock_p1, mock_p2]

    # Mock tables
    mock_table = MagicMock()
    mock_row1 = MagicMock()
    mock_cell1 = MagicMock()
    mock_cell1.text = "Cell 1\nwith space"
    mock_row1.cells = [mock_cell1]
    mock_table.rows = [mock_row1]
    mock_doc.tables = [mock_table]

    blocks = list(iter_block_text(mock_doc))

    assert len(blocks) == 3
    assert blocks[0] == "P|Paragraph 1"
    assert blocks[1] == "P|"
    assert blocks[2] == "T1R1C1|Cell 1 with space"


def test_load_docx_lines_validation_failure(tmp_path):
    non_existent = tmp_path / "no.docx"
    with pytest.raises(ValueError, match="File not found"):
        load_docx_lines(non_existent)


@patch("docxdiff.Document")
def test_load_docx_lines_success(mock_document_class, tmp_path):
    valid_file = tmp_path / "valid.docx"
    valid_file.write_text("fake zip")

    # Mock the Document object returned by Document(path)
    mock_doc = MagicMock()
    mock_p = MagicMock()
    mock_p.text = "Hello World"
    mock_doc.paragraphs = [mock_p]
    mock_doc.tables = []
    mock_document_class.return_value = mock_doc

    lines = load_docx_lines(valid_file)
    assert lines == ["P|Hello World"]


@patch("docxdiff.Document")
def test_load_docx_lines_runtime_error(mock_document_class, tmp_path):
    mock_document_class.side_effect = Exception("Corrupted file")
    valid_file = tmp_path / "corrupt.docx"
    valid_file.write_text("fake corrupt zip")

    with pytest.raises(RuntimeError, match="Failed to load document"):
        load_docx_lines(valid_file)


def test_structured_extraction_preserves_order_and_sections(tmp_path):
    path = tmp_path / "structured.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "Outer cell"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell"
    document.add_paragraph("After table")
    document.sections[0].header.paragraphs[0].text = "Header text"
    document.sections[0].footer.paragraphs[0].text = "Footer text"
    document.save(path)

    blocks = extract_document_blocks(path)

    assert [block.text for block in blocks] == [
        "Before table",
        "Outer cell",
        "Nested cell",
        "After table",
        "Header text",
        "Footer text",
    ]
    assert [block.kind for block in blocks] == [
        "paragraph",
        "table_cell",
        "table_cell",
        "paragraph",
        "header",
        "footer",
    ]
    assert blocks[2].location["table_path"] == [1, 1]


def test_structured_extraction_classifies_headings_and_lists(tmp_path):
    path = tmp_path / "styles.docx"
    document = Document()
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph("First item", style="List Bullet")
    document.save(path)

    blocks = extract_document_blocks(path)

    assert blocks[0].kind == "heading"
    assert blocks[0].style == "Heading 1"
    assert blocks[1].kind == "list_item"
    assert blocks[1].list_level == 0


def test_word_spans_handle_punctuation_unicode_and_normalization():
    old_spans, new_spans = compare_word_spans(
        "Hello, Wörld!  Value 10.",
        "Hello, wörld! Value 20.",
        ignore_case=True,
        ignore_whitespace=True,
    )

    assert [span.text for span in old_spans] == ["10"]
    assert [span.text for span in new_spans] == ["20"]


def test_compare_blocks_preserves_original_text_when_ignored():
    old = DocumentBlock(
        "a",
        "paragraph",
        "Important   TERMS",
        {"scope": "body", "paragraph": 1},
        normalized_text="important terms",
    )
    new = DocumentBlock(
        "b",
        "paragraph",
        "important terms",
        {"scope": "body", "paragraph": 1},
        normalized_text="important terms",
    )

    result = compare_blocks([old], [new], ignore_case=True, ignore_whitespace=True)

    assert result.differences_found is False
    assert result.changes[0].old_block.text == "Important   TERMS"
    assert result.changes[0].new_block.text == "important terms"


def test_compare_blocks_detects_moved_block():
    blocks_a = [
        DocumentBlock("a1", "paragraph", "Alpha", {"paragraph": 1}, normalized_text="Alpha"),
        DocumentBlock("a2", "paragraph", "Beta", {"paragraph": 2}, normalized_text="Beta"),
        DocumentBlock("a3", "paragraph", "Gamma", {"paragraph": 3}, normalized_text="Gamma"),
    ]
    blocks_b = [
        DocumentBlock("b1", "paragraph", "Beta", {"paragraph": 1}, normalized_text="Beta"),
        DocumentBlock("b2", "paragraph", "Alpha", {"paragraph": 2}, normalized_text="Alpha"),
        DocumentBlock("b3", "paragraph", "Gamma", {"paragraph": 3}, normalized_text="Gamma"),
    ]

    result = compare_blocks(blocks_a, blocks_b)

    moved = [change for change in result.changes if change.change_type == "moved"]
    assert len(moved) == 1
    assert moved[0].old_block.text == "Beta"
    assert moved[0].new_block.text == "Beta"
    assert moved[0].moved_from == {"paragraph": 2}
    assert moved[0].moved_to == {"paragraph": 1}


def test_compare_documents_can_be_cancelled(tmp_path):
    path = tmp_path / "cancel.docx"
    document = Document()
    document.add_paragraph("Text")
    document.save(path)
    cancel_event = threading.Event()
    cancel_event.set()

    with pytest.raises(ComparisonCancelled):
        compare_documents(path, path, cancel_event=cancel_event)


def test_print_unified_diff(capsys):
    a_lines = ["line 1", "line 2"]
    b_lines = ["line 1", "line 2 changed"]

    any_diff = print_unified_diff(a_lines, b_lines, "file_a.docx", "file_b.docx")
    assert any_diff

    captured = capsys.readouterr()
    assert "--- file_a.docx" in captured.out
    assert "+++ file_b.docx" in captured.out
    assert "-line 2" in captured.out
    assert "+line 2 changed" in captured.out


def test_print_colored_diff(capsys):
    a_lines = ["line 1", "line 2"]
    b_lines = ["line 1", "line 2 changed"]

    any_diff = print_colored_diff(a_lines, b_lines, "file_a.docx", "file_b.docx")
    assert any_diff

    captured = capsys.readouterr()
    # Check that ANSI color codes or correct diff lines are printed
    assert "file_a.docx" in captured.out
    assert "file_b.docx" in captured.out
    assert "line 2" in captured.out


def test_calculate_diff_stats():
    a_lines = ["line 1", "line 2", "line 3"]
    b_lines = ["line 1", "line 2 changed", "line 4"]

    stats = calculate_diff_stats(a_lines, b_lines)
    assert stats["additions"] == 2
    assert stats["deletions"] == 2
    assert stats["unchanged"] == 1
    assert stats["total_lines_a"] == 3
    assert stats["total_lines_b"] == 3
    assert isinstance(stats["similarity"], float)


def test_export_to_html(tmp_path):
    a_lines = ["line 1"]
    b_lines = ["line 2"]
    out_html = tmp_path / "diff.html"

    success = export_to_html(a_lines, b_lines, "a.docx", "b.docx", str(out_html))
    assert success
    assert out_html.exists()
    content = out_html.read_text(encoding="utf-8")
    assert "DOCX Diff v" in content
    assert "a.docx" in content
    assert "b.docx" in content


def test_export_to_html_escapes_descriptions(tmp_path):
    out_html = tmp_path / "diff.html"

    export_to_html(["a"], ["b"], "<script>alert(1)</script>", "b.docx", str(out_html))

    content = out_html.read_text(encoding="utf-8")
    assert "<script>alert(1)</script>" not in content
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in content


def test_export_to_json(tmp_path):
    a_lines = ["line 1"]
    b_lines = ["line 2"]
    out_json = tmp_path / "diff.json"

    success = export_to_json(a_lines, b_lines, "a.docx", "b.docx", str(out_json))
    assert success
    assert out_json.exists()

    data = json.loads(out_json.read_text(encoding="utf-8"))
    assert "metadata" in data
    assert "statistics" in data
    assert "changes" in data
    assert data["statistics"]["similarity"] == 0.0


def test_export_to_json_includes_provenance(tmp_path):
    out_json = tmp_path / "diff.json"
    sources = {
        "file_a": {"path": "a.docx", "size_bytes": 1, "sha256": "a" * 64},
        "file_b": {"path": "b.docx", "size_bytes": 1, "sha256": "b" * 64},
    }

    export_to_json(
        ["line 1"],
        ["line 2"],
        "a.docx",
        "b.docx",
        str(out_json),
        source_metadata=sources,
        generated_at="2026-06-07T12:00:00+00:00",
        report_id="report-123",
    )

    metadata = json.loads(out_json.read_text(encoding="utf-8"))["metadata"]
    assert metadata["sources"] == sources
    assert metadata["timestamp"] == "2026-06-07T12:00:00+00:00"
    assert metadata["report_id"] == "report-123"


def test_structured_exports_share_word_level_result(tmp_path):
    file_a = tmp_path / "a.docx"
    file_b = tmp_path / "b.docx"
    html_report = tmp_path / "report.html"
    json_report = tmp_path / "report.json"
    document_a = Document()
    document_a.add_paragraph("Payment is due in 30 days.")
    document_a.save(file_a)
    document_b = Document()
    document_b.add_paragraph("Payment is due in 45 calendar days.")
    document_b.save(file_b)
    result = compare_documents(file_a, file_b)

    export_to_html(
        result.legacy_lines_a(),
        result.legacy_lines_b(),
        str(file_a),
        str(file_b),
        str(html_report),
        structured_result=result,
    )
    export_to_json(
        result.legacy_lines_a(),
        result.legacy_lines_b(),
        str(file_a),
        str(file_b),
        str(json_report),
        structured_result=result,
    )

    html_content = html_report.read_text(encoding="utf-8")
    data = json.loads(json_report.read_text(encoding="utf-8"))
    assert data["schema_version"] == 2
    assert "changes" in data
    structured_change = data["structured_comparison"]["changes"][0]
    assert structured_change["old_spans"][0]["text"] == "30"
    assert "word-del" in html_content
    assert ">30<" in html_content


def test_print_side_by_side(capsys):
    a_lines = ["line 1", "line 2"]
    b_lines = ["line 1", "line 3"]

    print_side_by_side(a_lines, b_lines, "a.docx", "b.docx", width=40)
    captured = capsys.readouterr()
    assert "a.docx" in captured.out
    assert "b.docx" in captured.out
    assert "line 1" in captured.out


def test_print_statistics(capsys):
    stats = {
        "similarity": 80.0,
        "additions": 1,
        "deletions": 1,
        "unchanged": 4,
        "total_lines_a": 5,
        "total_lines_b": 5,
    }
    print_statistics(stats, verbose=True)
    captured = capsys.readouterr()
    assert "DIFF STATISTICS" in captured.out
    assert "Similarity:        80.0%" in captured.out
    assert "Change Breakdown:" in captured.out


def test_main_cli_success(capsys, tmp_path):
    file_a = tmp_path / "a.docx"
    file_b = tmp_path / "b.docx"
    document_a = Document()
    document_a.add_paragraph("line 1")
    document_a.add_paragraph("line 2")
    document_a.save(file_a)
    document_b = Document()
    document_b.add_paragraph("line 1")
    document_b.add_paragraph("line 2 changed")
    document_b.save(file_b)

    test_args = ["docxdiff.py", str(file_a), str(file_b)]
    with patch.object(sys, "argv", test_args):
        exit_code = main()
        assert exit_code == 1  # Differences found

    captured = capsys.readouterr()
    assert "line 2" in captured.out


def test_main_ignore_options_keep_original_text_but_ignore_difference(capsys, tmp_path):
    file_a = tmp_path / "a.docx"
    file_b = tmp_path / "b.docx"
    document_a = Document()
    document_a.add_paragraph("Important   TERMS")
    document_a.save(file_a)
    document_b = Document()
    document_b.add_paragraph("important terms")
    document_b.save(file_b)

    test_args = [
        "docxdiff.py",
        str(file_a),
        str(file_b),
        "--ignore-case",
        "--ignore-whitespace",
    ]
    with patch.object(sys, "argv", test_args):
        assert main() == 0

    assert "No differences found." in capsys.readouterr().out


@patch("docxdiff.load_docx_lines")
def test_main_cli_same_files(mock_load, capsys, tmp_path):
    file_a = tmp_path / "a.docx"
    file_a.write_text("a")

    test_args = ["docxdiff.py", str(file_a), str(file_a)]
    with patch.object(sys, "argv", test_args):
        exit_code = main()
        assert exit_code == 2

    captured = capsys.readouterr()
    assert "Error: Both files point to the same location" in captured.err


def test_main_policy_violation_writes_redacted_audit_log(capsys, tmp_path):
    file_a = tmp_path / "sensitive-a.docx"
    file_b = tmp_path / "sensitive-b.docx"
    audit_log = tmp_path / "audit.jsonl"
    document_a = Document()
    document_a.add_paragraph("original")
    document_a.save(file_a)
    document_b = Document()
    document_b.add_paragraph("replacement")
    document_b.save(file_b)

    test_args = [
        "docxdiff.py",
        str(file_a),
        str(file_b),
        "--quiet",
        "--min-similarity",
        "90",
        "--audit-log",
        str(audit_log),
        "--audit-redact-paths",
    ]
    with patch.object(sys, "argv", test_args):
        assert main() == 3

    captured = capsys.readouterr()
    assert captured.out == ""
    event = json.loads(audit_log.read_text(encoding="utf-8").strip())
    assert event["schema_version"] == 1
    assert event["sources"]["file_a"]["path"] == file_a.name
    assert event["sources"]["file_b"]["path"] == file_b.name
    assert len(event["sources"]["file_a"]["sha256"]) == 64
    assert event["result"]["exit_code"] == 3
    assert event["result"]["policy_violations"]


@patch("docxdiff.load_docx_lines")
def test_main_rejects_output_overwriting_input(mock_load, capsys, tmp_path):
    file_a = tmp_path / "a.docx"
    file_b = tmp_path / "b.docx"
    file_a.write_bytes(b"a")
    file_b.write_bytes(b"b")

    test_args = ["docxdiff.py", str(file_a), str(file_b), "--json", str(file_a)]
    with patch.object(sys, "argv", test_args):
        assert main() == 2

    mock_load.assert_not_called()
    assert "overwrite an input file" in capsys.readouterr().err


def test_main_end_to_end_enterprise_reports(tmp_path):
    file_a = tmp_path / "a.docx"
    file_b = tmp_path / "b.docx"
    json_report = tmp_path / "report.json"
    audit_log = tmp_path / "audit.jsonl"

    document_a = Document()
    document_a.add_paragraph("Approved contract")
    document_a.save(file_a)

    document_b = Document()
    document_b.add_paragraph("Revised contract")
    document_b.save(file_b)

    test_args = [
        "docxdiff.py",
        str(file_a),
        str(file_b),
        "--quiet",
        "--json",
        str(json_report),
        "--audit-log",
        str(audit_log),
    ]
    with patch.object(sys, "argv", test_args):
        assert main() == 1

    report = json.loads(json_report.read_text(encoding="utf-8"))
    event = json.loads(audit_log.read_text(encoding="utf-8").strip())
    assert report["metadata"]["sources"]["file_a"]["sha256"] == sha256_file(file_a)
    assert report["metadata"]["report_id"] == event["event_id"]
    assert event["result"]["differences_found"] is True


def test_display_side_by_side_alignment():
    # Create a mock instance of DocxDiffGUI without calling __init__
    gui = MagicMock(spec=DocxDiffGUI)

    # Mock text widgets and their insert methods
    gui.left_text = MagicMock()
    gui.right_text = MagicMock()

    # Mock show_line_numbers_var
    gui.show_line_numbers_var = MagicMock()
    gui.show_line_numbers_var.get.return_value = True

    # Prepare inputs: 'replace' scenario
    # a_lines has 2 lines, b_lines has 1 line
    a_lines = ["old 1", "old 2"]
    b_lines = ["new 1"]

    matcher = difflib.SequenceMatcher(None, a_lines, b_lines)

    # Call display_side_by_side using the mocked gui instance
    DocxDiffGUI.display_side_by_side(gui, a_lines, b_lines, matcher)

    # Check that delete was called on both text widgets
    gui.left_text.delete.assert_any_call(1.0, tk.END)
    gui.right_text.delete.assert_any_call(1.0, tk.END)

    # In the replace block, 2 lines are removed from left, and 1 line is added to right
    # Because of our alignment fix, 1 empty spacer line should have been added to the right side!
    # Let's count how many times "     │ \n" was inserted into the right text widget.
    # It should be exactly 1 time!
    insert_calls = gui.right_text.insert.call_args_list
    spacer_calls = [c for c in insert_calls if c[0][1] == "     │ \n"]
    assert len(spacer_calls) == 1


def test_gui_swap_files():
    # Mock DocxDiffGUI instance
    gui = MagicMock(spec=DocxDiffGUI)
    gui.file_a_var = MagicMock()
    gui.file_b_var = MagicMock()
    gui.status_bar = MagicMock()

    gui.file_a_var.get.return_value = "path/to/a.docx"
    gui.file_b_var.get.return_value = "path/to/b.docx"

    DocxDiffGUI.swap_files(gui)

    gui.file_a_var.set.assert_called_once_with("path/to/b.docx")
    gui.file_b_var.set.assert_called_once_with("path/to/a.docx")
    gui.status_bar.config.assert_called_once_with(text="Swapped File A and File B")


def test_gui_create_menu_bar():
    class FakeMenu:
        def __init__(self, parent, tearoff=0):
            self.parent = parent
            self.tearoff = tearoff
            self.items = []

        def add_command(self, **kwargs):
            self.items.append(("command", kwargs))

        def add_separator(self):
            self.items.append(("separator", {}))

        def add_cascade(self, **kwargs):
            self.items.append(("cascade", kwargs))

        def add_checkbutton(self, **kwargs):
            self.items.append(("checkbutton", kwargs))

    gui = MagicMock(spec=DocxDiffGUI)
    gui.root = MagicMock()
    gui.ignore_case_var = MagicMock()
    gui.ignore_whitespace_var = MagicMock()
    gui.filter_var = MagicMock()
    gui.show_line_numbers_var = MagicMock()
    gui.sync_scroll_var = MagicMock()
    gui.results_text = MagicMock()

    with patch("docxdiff.tk.Menu", side_effect=FakeMenu):
        DocxDiffGUI.create_menu_bar(gui)

    top_labels = [
        item[1]["label"] for item in gui.menu_bar.items if item[0] == "cascade"
    ]
    assert top_labels == ["File", "Edit", "Compare", "View", "Help"]
    gui.root.config.assert_called_once_with(menu=gui.menu_bar)


@pytest.mark.parametrize(
    ("width", "expected"),
    [
        (2400, "wide"),
        (1400, "medium"),
        (900, "compact"),
        (640, "narrow"),
    ],
)
def test_gui_responsive_mode_breakpoints(width, expected):
    assert DocxDiffGUI.responsive_mode_for_width(width) == expected


def test_gui_reset_options():
    gui = MagicMock(spec=DocxDiffGUI)
    gui.context_var = MagicMock()
    gui.ignore_case_var = MagicMock()
    gui.ignore_whitespace_var = MagicMock()
    gui.show_line_numbers_var = MagicMock()
    gui.filter_var = MagicMock()
    gui.sync_scroll_var = MagicMock()
    gui.status_bar = MagicMock()

    DocxDiffGUI.reset_options(gui)

    gui.context_var.set.assert_called_once_with(Config.DEFAULT_CONTEXT_LINES)
    gui.ignore_case_var.set.assert_called_once_with(False)
    gui.ignore_whitespace_var.set.assert_called_once_with(False)
    gui.show_line_numbers_var.set.assert_called_once_with(False)
    gui.filter_var.set.assert_called_once_with(False)
    gui.sync_scroll_var.set.assert_called_once_with(True)
    gui.reset_zoom.assert_called_once_with(update_status=False)


def test_gui_new_comparison():
    gui = MagicMock(spec=DocxDiffGUI)
    gui.file_a_var = MagicMock()
    gui.file_b_var = MagicMock()
    gui.file_a_combo = MagicMock()
    gui.status_bar = MagicMock()

    DocxDiffGUI.new_comparison(gui)

    gui.file_a_var.set.assert_called_once_with("")
    gui.file_b_var.set.assert_called_once_with("")
    gui.reset_options.assert_called_once_with()
    gui.clear_results.assert_called_once_with()
    gui.file_a_combo.focus_set.assert_called_once_with()
    gui.status_bar.config.assert_called_once_with(text="Ready for a new comparison")


def test_gui_comparison_starts_without_blocking(tmp_path):
    file_a = tmp_path / "a.docx"
    file_b = tmp_path / "b.docx"
    file_a.write_bytes(b"a")
    file_b.write_bytes(b"b")
    release = threading.Event()
    result = compare_blocks(
        [DocumentBlock("a", "paragraph", "A", {}, normalized_text="A")],
        [DocumentBlock("b", "paragraph", "B", {}, normalized_text="B")],
    )

    gui = MagicMock(spec=DocxDiffGUI)
    gui.file_a_var = MagicMock()
    gui.file_b_var = MagicMock()
    gui.file_a_var.get.return_value = str(file_a)
    gui.file_b_var.get.return_value = str(file_b)
    gui.ignore_case_var = MagicMock()
    gui.ignore_whitespace_var = MagicMock()
    gui.ignore_case_var.get.return_value = False
    gui.ignore_whitespace_var.get.return_value = False
    gui.compare_button = MagicMock()
    gui.cancel_button = MagicMock()
    gui.status_bar = MagicMock()
    gui.root = MagicMock()
    gui._comparison_thread = None
    gui._comparison_cancel = threading.Event()

    def slow_worker(*args, **kwargs):
        release.wait(timeout=2)
        gui.current_comparison_result = result

    gui._comparison_worker = slow_worker
    started = time.perf_counter()
    DocxDiffGUI.compare_files(gui)
    elapsed = time.perf_counter() - started
    assert elapsed < 0.2
    assert gui._comparison_thread.is_alive()
    release.set()
    gui._comparison_thread.join(timeout=2)

    gui.compare_button.config.assert_called_with(state=tk.DISABLED)
    gui.cancel_button.config.assert_called_with(state=tk.NORMAL)


def test_gui_cancel_comparison_sets_event():
    gui = MagicMock(spec=DocxDiffGUI)
    gui._comparison_thread = MagicMock()
    gui._comparison_thread.is_alive.return_value = True
    gui._comparison_cancel = threading.Event()
    gui.status_bar = MagicMock()

    DocxDiffGUI.cancel_comparison(gui)

    assert gui._comparison_cancel.is_set()
    gui.status_bar.config.assert_called_once_with(text="Cancelling comparison...")


def test_gui_worker_delivers_result_through_queue(tmp_path):
    file_a = tmp_path / "a.docx"
    file_b = tmp_path / "b.docx"
    result = compare_blocks(
        [DocumentBlock("a", "paragraph", "A", {})],
        [DocumentBlock("b", "paragraph", "B", {})],
    )
    gui = MagicMock(spec=DocxDiffGUI)
    gui._comparison_cancel = threading.Event()
    gui._comparison_queue = queue.Queue()
    gui.root = MagicMock()

    with patch("docxdiff.compare_documents", return_value=result):
        DocxDiffGUI._comparison_worker(
            gui,
            file_a,
            file_b,
            {"ignore_case": False, "ignore_whitespace": False},
        )

    message, payload = gui._comparison_queue.get_nowait()
    assert message == "complete"
    assert payload == (file_a, file_b, result)
    gui.root.after.assert_not_called()


def test_gui_add_to_history():
    gui = MagicMock(spec=DocxDiffGUI)
    gui.history = []
    gui.save_history = MagicMock()
    gui.update_history_comboboxes = MagicMock()

    # Test adding a path
    DocxDiffGUI.add_to_history(gui, "test1.docx")
    assert str(Path("test1.docx").resolve()) in gui.history

    # Test limiting history
    gui.history = [f"file_{i}.docx" for i in range(15)]
    DocxDiffGUI.add_to_history(gui, "new.docx")
    assert len(gui.history) == 10
    assert str(Path("new.docx").resolve()) in gui.history


@patch("docxdiff.messagebox")
@patch("docxdiff.filedialog.asksaveasfilename")
@patch("docxdiff.export_to_json")
def test_gui_export_json_success(mock_export, mock_ask, mock_msg):
    gui = MagicMock(spec=DocxDiffGUI)
    gui.current_diff_lines = ["some diff"]
    old = DocumentBlock("old", "paragraph", "line 1", {"scope": "body"})
    new = DocumentBlock("new", "paragraph", "line 2", {"scope": "body"})
    gui.current_comparison_result = compare_blocks([old], [new])
    gui.file_a_var = MagicMock()
    gui.file_b_var = MagicMock()
    gui.file_a_var.get.return_value = "a.docx"
    gui.file_b_var.get.return_value = "b.docx"
    gui.ignore_whitespace_var = MagicMock()
    gui.ignore_whitespace_var.get.return_value = False
    gui.ignore_case_var = MagicMock()
    gui.ignore_case_var.get.return_value = False
    gui.context_var = MagicMock()
    gui.context_var.get.return_value = 3
    gui.status_bar = MagicMock()

    mock_ask.return_value = "output.json"

    DocxDiffGUI.export_json(gui)

    mock_export.assert_called_once()
    gui.status_bar.config.assert_called_once_with(text="Exported to output.json")
    mock_msg.showinfo.assert_called_once_with(
        "Success", "Results exported to output.json"
    )
